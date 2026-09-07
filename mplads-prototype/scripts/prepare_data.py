"""Read the workbook into a local-only cache outside the application."""
import argparse, hashlib, json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT.parent / 'outputs/01a06d6b-9b94-7a61-b28b-6f9116f20942/MPLADS_Connected_Feature_Dataset_2026-09-05.xlsx'
def clean(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (datetime, date)) else v
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', type=Path, default=DEFAULT)
    args = parser.parse_args()
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    tables = {}
    for name in ['Work_Features', 'MP_Features', 'IDA_Features', 'Calamity_Features', 'Duplicate_Candidates', 'Feature_Dictionary']:
        iterator = book[name].iter_rows(min_row=4, values_only=True)
        columns = list(next(iterator))
        rows = [[clean(v) for v in row[:len(columns)]] for row in iterator if row[0] is not None]
        tables[name] = {'columns': columns, 'rows': rows}
    work = [dict(zip(tables['Work_Features']['columns'], r)) for r in tables['Work_Features']['rows']]
    ids = {r['work_id'] for r in work}
    pairs = [dict(zip(tables['Duplicate_Candidates']['columns'], r)) for r in tables['Duplicate_Candidates']['rows']]
    assert len(work) == len(ids) == 10000
    assert len(pairs) == 5106
    assert all(p['work_id_a'] in ids and p['work_id_b'] in ids and p['work_id_a'] != p['work_id_b'] for p in pairs)
    assert len({tuple(sorted((p['work_id_a'], p['work_id_b']))) for p in pairs}) == len(pairs)
    assert abs(sum(r['sanction_amount'] for r in work) - 5276800278) < 0.01
    assert all(r['allocation_join_status'] == 'Matched' for r in work)
    bands = dict(Counter(r['review_priority_band'] for r in work))
    assert bands == {'Routine': 6820, 'Low': 2626, 'Medium': 544, 'High': 10}, bands
    rules = [list(row[:5]) for row in book['Scoring_Rules'].iter_rows(min_row=5, values_only=True) if row[0] is not None]
    payload = {'meta': {'snapshot': '2026-09-01', 'sourceFile': args.workbook.name, 'sourceSha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(), 'reportedWorksTotal': 41388495369.08, 'visibleWorksTotal': 5276800278, 'allocationTotal': 83336673298.01, 'calamityTotal': 40567400, 'bands': bands, 'guidelines': 'https://www.mplads.gov.in/MPLADS/UploadedFiles/MPLADSGuidelinesApril2023.pdf', 'counts': {name: len(t['rows']) for name, t in tables.items()}}, 'tables': tables, 'rules': rules}
    destination = ROOT.parent / 'prototype-local-data/snapshot.json'
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':'), allow_nan=False), encoding='utf-8')
    book.close()
    print(json.dumps({'counts': payload['meta']['counts'], 'bands': bands, 'bytes': destination.stat().st_size, 'sourceSha256': payload['meta']['sourceSha256']}))
if __name__ == '__main__':
    main()
