"""Read-only XLSX verification; writes only a separate verification JSON."""
from datetime import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'data_preparation/local/release_2026-09-10_v2'
FILE=ROOT/'outputs/ps26102-all-cohorts-20260910/MPLADS_Data_Review.xlsx'

def run():
    expected=json.loads((DATA/'review_data.json').read_text(encoding='utf-8'))
    formulas=load_workbook(FILE,data_only=False)
    values=load_workbook(FILE,data_only=True)
    checks=[];cells=0
    def require(condition,message):
        if not condition:raise ValueError(message)
        checks.append(dict(check=message,passed=True))
    def matches(actual,value,key):
        if value is None:return actual is None
        if key.endswith('_date') and isinstance(value,str):return isinstance(actual,datetime) and actual.date().isoformat()==value
        if key.endswith('_inr'):return Decimal(str(actual)).quantize(Decimal('.01'))==Decimal(str(value)).quantize(Decimal('.01'))
        if isinstance(value,str) and value.lstrip().startswith(('=','+','@','-')):return actual in (value,"'"+value)
        return actual==value
    specs=[
        ('Overview',6,[dict(metric=a,value=b) for a,b in expected['overview']],['metric','value']),
        ('Overview',20,expected['signals'],['signal','works','interpretation']),
        ('MP terms',6,expected['mp_terms'],['mp_name','cohort','allocation_state','tenure_start_date','tenure_end_date','allocated_inr','work_count','sanctioned_count','completed_count','successful_payment_inr','pending_payment_inr','open_over_one_year_count','mp_key']),
        ('Review sample',6,expected['review_sample'],['work_id','cohort','mp_name','state','activity_type','screening_signal_count','sanction_amount_inr','successful_payment_inr','screening_reasons']),
        ('Source audit',6,expected['sources'],['cohort','kind','source_rows','accepted_rows','quarantined_rows','column_count','source_file']),
        ('Source audit',28,expected['cohorts'],['cohort','work_count','recommended_record_count','sanctioned_count','completed_count','successful_payment_inr','pending_payment_inr']),
        ('Feature dictionary',6,expected['dictionary'],['table','field','unit','definition','model_use']),
        ('Limits and sources',6,expected['limits'],['topic','limitation']),
    ]
    for sheet,start,rows,columns in specs:
        s=values[sheet]
        for i,row in enumerate(rows,start):
            for j,key in enumerate(columns,1):
                actual=s.cell(i,j).value;target=row[key]
                if not matches(actual,target,key):raise ValueError(f'{sheet}!{s.cell(i,j).coordinate} differs from prepared data')
                cells+=1
        require(True,f'{sheet}, row {start}: {len(rows)} rows match prepared data')
    require(values.sheetnames==['Overview','MP terms','Review sample','Source audit','Feature dictionary','Limits and sources'],'All six intended sheets; no empty default sheet')
    for s in formulas:
        require(not s.sheet_view.showGridLines,f'{s.title}: gridlines hidden')
        for t in s.tables.values():require(t.autoFilter is not None,f'{t.name}: filter is saved')
    for sheet,pane in [('MP terms','B6'),('Review sample','B6'),('Feature dictionary','C6'),('Source audit','A6')]:require(formulas[sheet].freeze_panes==pane,f'{sheet}: frozen identifiers/header retained')
    require(isinstance(values['MP terms']['D6'].value,datetime),'Dates are stored as sortable dates')
    require(formulas['MP terms']['F6'].number_format=='#,##0.00','INR amounts display two decimal places')
    for i,key in enumerate(['work_count','recommended_record_count','sanctioned_count','completed_count','successful_payment_inr','pending_payment_inr'],2):
        expected_sum=sum(row[key] for row in expected['cohorts']);actual=values['Source audit'].cell(31,i).value
        require(matches(actual,expected_sum,key),f'Cohort total {key} has correct cached result')
        require(formulas['Source audit'].cell(31,i).data_type=='f',f'Cohort total {key} remains a formula')
    errors=[f'{s.title}!{c.coordinate}' for s in values for row in s for c in row if c.data_type=='e']
    require(not errors,'No saved formula errors')
    # Source text must not become executable formulas in the review data.
    unexpected=[f'{s.title}!{c.coordinate}' for s in formulas for row in s for c in row if c.data_type=='f' and not(s.title=='Source audit' and c.row==31 and 2<=c.column<=7)]
    require(not unexpected,'No source text converted to spreadsheet formulas')
    report=dict(workbook=str(FILE),sha256=hashlib.sha256(FILE.read_bytes()).hexdigest(),data_cells_compared=cells,checks_passed=len(checks),checks=checks,native_excel_ui_tested=False,visual_review='See workbook_verification.json and previews; this checker does not replace visual review.')
    (DATA/'workbook_saved_file_verification.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    formulas.close();values.close()
    print(json.dumps({k:v for k,v in report.items() if k!='checks'},indent=2))

if __name__=='__main__':run()
