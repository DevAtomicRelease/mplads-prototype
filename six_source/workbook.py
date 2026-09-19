"""Plain Excel review workbook for officers, built from the local six_source outputs.

Reads local/ (CSV outputs + audit.json + ab_metrics.json) and writes
local/MPLADS_Review.xlsx with human-readable sheets: overview, priority cases,
MP / authority / vendor summaries, the A/B result, the feature dictionary and the
sources & limits. Amounts are shown in INR. Run after build.py (and validate.py).

    python workbook.py
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
HEAD = Font(bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="10283F")
INR_COLS = ("_paise",)


def _sheet(wb, title, frame, money_to_inr=True, width=22, wrap=None):
    ws = wb.create_sheet(title[:31])
    frame = frame.copy()
    money = [c for c in frame.columns if money_to_inr and c.endswith("_paise")]
    for c in money:
        frame[c] = (pd.to_numeric(frame[c], errors="coerce") / 100).round(2)  # paise -> rupees
    # A converted column holds rupees now; relabel so no rupee value keeps a paise name.
    frame = frame.rename(columns={c: c[:-len("_paise")].rstrip("_").replace("_", " ") + " (INR)" for c in money})
    ws.append(list(frame.columns))
    for cell in ws[1]:
        cell.font = HEAD; cell.fill = HEADFILL; cell.alignment = Alignment(vertical="center")
    for row in frame.itertuples(index=False, name=None):
        ws.append(["" if pd.isna(v) else v for v in row])
    ws.freeze_panes = "A2"
    for i, c in enumerate(frame.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, (wrap or {}).get(c, width)))
    return ws


def build(local: Path, out: Path | None = None):
    out = out or local / "MPLADS_Review.xlsx"
    meta = json.loads((local / "audit.json").read_text(encoding="utf-8"))
    rd = lambda name: pd.read_csv(local / name, dtype=str, keep_default_na=False)
    wb = Workbook(); wb.remove(wb.active)

    # Overview
    t = meta["totals"]; inr = lambda p: round(p / 1e9, 2)  # paise -> INR crore (1 cr = 1e7 rupees = 1e9 paise)
    ov = [("Snapshot date", meta["as_of"]), ("Cohorts", ", ".join(meta.get("cohorts", []))),
          ("Works", t["works"]), ("Recommended", t["recommendations"]), ("Sanctioned", t["sanctions"]),
          ("Reported complete", t["completions"]),
          ("Sanctioned amount (INR cr)", inr(t["sanction_paise"])), ("Settled reported (INR cr)", inr(t["successful_payment_paise"])),
          ("In-progress (INR cr)", inr(t["pending_payment_paise"]))]
    ov += [(f"Rule: {k}", v) for k, v in meta["rule_counts"].items()]
    ov += [(f"Band: {k}", v) for k, v in meta.get("priority_counts", {}).items()]
    ws = wb.create_sheet("Overview")
    ws.append(["MPLADS review workbook — signals for verification, not findings of fraud"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for k, v in ov:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 30

    # Priority cases
    work = rd("Work_Features.csv")
    keep = [c for c in ["work_id","cohort","mp_name","state","ida_name","activity_type","description","lifecycle",
                        "sanction_date","sanction_amount_paise","successful_payment_paise","priority_score","priority_band",
                        "reason_codes","isolation_percentile","dbscan_outlier_flag"] if c in work.columns]
    top = work.assign(_s=pd.to_numeric(work.priority_score, errors="coerce")).sort_values(["_s","work_id"], ascending=[False, True]).head(1000)[keep]
    _sheet(wb, "Priority cases", top, wrap={"description": 60, "reason_codes": 40, "ida_name": 34})

    for name, title, sort, n in [("MP_Features.csv","MP summary","successful_payment_paise",1023),
                                 ("IDA_Features.csv","Authority summary","successful_payment_paise",1045),
                                 ("Vendor_Features.csv","Vendor summary","successful_payment_paise",500)]:
        if (local / name).is_file():
            f = rd(name)
            if sort in f.columns:
                f = f.assign(_s=pd.to_numeric(f[sort], errors="coerce")).sort_values("_s", ascending=False).drop(columns="_s").head(n)
            _sheet(wb, title, f)

    # A/B result
    ab = local / "ab_metrics.json"
    if ab.is_file():
        m = json.loads(ab.read_text(encoding="utf-8"))
        rows = [{"budget": f"{b['fraction']*100:.0f}%", "A recovery": round(b["a_recovery"], 3), "B recovery": round(b["b_recovery"], 3),
                 "B-A pp": round(b["difference"] * 100, 1), "95% low pp": round(b["ci95"][0] * 100, 1), "95% high pp": round(b["ci95"][1] * 100, 1),
                 "real queue overlap": round(b.get("actual_queue_overlap", 0), 3)} for b in m["budgets"]]
        _sheet(wb, "A-B validation", pd.DataFrame(rows), money_to_inr=False, width=16)

    if (local / "Feature_Dictionary.csv").is_file():
        _sheet(wb, "Feature dictionary", rd("Feature_Dictionary.csv"), money_to_inr=False, wrap={"definition": 60, "field": 34})

    src = pd.DataFrame([{"file": s["file"], "rows": s["rows"], "accepted": s.get("accepted_rows"), "quarantined": s.get("quarantined_rows"), "sha256": s["sha256"]} for s in meta["sources"]])
    ws = _sheet(wb, "Sources & limits", src, money_to_inr=False, wrap={"file": 34, "sha256": 66})
    ws.append([]); ws.append(["Evidence limits:"])
    for lim in meta.get("limits", []):
        ws.append([lim])

    wb.save(out)
    print(f"Wrote {out} ({len(wb.sheetnames)} sheets)")
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--local", type=Path, default=HERE / "local")
    args = parser.parse_args()
    build(args.local)
